"""
Capa de Servicio - Generador de Excel de Resultados
Arma en memoria un archivo .xlsx con el resumen, los estudiantes procesados
y las inconsistencias detectadas. No escribe nada en el disco.
"""

import io
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Colores institucionales usados en los encabezados del Excel.
COLOR_AZUL_MARINO = "14213D"
COLOR_BLANCO = "FFFFFF"


def _estilo_encabezado(hoja, encabezados: List[str]) -> None:
    """
    Escribe la primera fila con los títulos y le aplica el estilo institucional.
    """
    hoja.append(encabezados)

    for numero_columna, _ in enumerate(encabezados, start=1):
        celda = hoja.cell(row=1, column=numero_columna)
        celda.font = Font(bold=True, color=COLOR_BLANCO)
        celda.fill = PatternFill("solid", fgColor=COLOR_AZUL_MARINO)
        celda.alignment = Alignment(horizontal="center", vertical="center")

    # La fila de títulos queda fija al desplazarse por el Excel.
    hoja.freeze_panes = "A2"
    hoja.row_dimensions[1].height = 22


def _ajustar_anchos(hoja) -> None:
    """
    Calcula un ancho razonable para cada columna según su contenido.
    """
    for columna in hoja.columns:
        largo_maximo = 0
        for celda in columna:
            if celda.value is not None:
                largo_maximo = max(largo_maximo, len(str(celda.value)))

        letra = get_column_letter(columna[0].column)
        hoja.column_dimensions[letra].width = min(largo_maximo + 2, 45)


def _escribir_tabla(hoja, encabezados: List[str], filas: List[List[Any]]) -> None:
    """
    Escribe una tabla completa: encabezados, filas y ajuste de anchos.
    """
    _estilo_encabezado(hoja, encabezados)

    for fila in filas:
        hoja.append(fila)

    if filas:
        hoja.auto_filter.ref = hoja.dimensions

    _ajustar_anchos(hoja)


def generar_excel_resultados(
    datos: Dict[str, Any], origen_datos: str
) -> io.BytesIO:
    """
    Genera el Excel de resultados y lo devuelve como un buffer en memoria.

    BytesIO es como una caja de archivo dentro de la memoria RAM: guarda los
    bytes igual que una carpeta guarda hojas, pero sin usar el disco. Al
    terminar, Flask puede enviar esa caja directamente al navegador.
    """
    libro = Workbook()

    # ----------------------------- Hoja Resumen -----------------------------
    resumen = libro.active
    resumen.title = "Resumen"

    resumen.append(["Indicador", "Valor"])
    resumen.append(["Origen de los datos", origen_datos])
    resumen.append(["Aprobación", datos["resumen"]["aprobacion"]])
    resumen.append(["Participación", datos["resumen"]["participacion"]])
    resumen.append(["Sin certificado", datos["resumen"]["sin_certificado"]])
    resumen.append(["Inconsistencias", datos["resumen"]["inconsistencias"]])
    resumen.append(["Total de estudiantes", datos["resumen"]["total_estudiantes"]])
    resumen.append(["Total de evaluaciones", datos["resumen"]["total_evaluaciones"]])
    resumen.append(["Total certificables", datos["resumen"]["total_certificables"]])

    for numero_columna in (1, 2):
        celda = resumen.cell(row=1, column=numero_columna)
        celda.font = Font(bold=True, color=COLOR_BLANCO)
        celda.fill = PatternFill("solid", fgColor=COLOR_AZUL_MARINO)
        celda.alignment = Alignment(horizontal="center", vertical="center")

    _ajustar_anchos(resumen)

    # --------------------------- Hoja Resultados ----------------------------
    resultados = libro.create_sheet("Resultados")
    encabezados_resultados = [
        "Identificación",
        "Nombre completo",
        "Correo",
        "Programa",
        "Cohorte",
        "Módulos",
        "Promedio nota",
        "Asistencia (%)",
        "Estado",
        "Certificable",
        "Motivo de inconsistencia",
    ]

    filas_resultados = []
    for estudiante in datos["estudiantes"]:
        filas_resultados.append(
            [
                estudiante.get("identificacion"),
                estudiante.get("nombre_completo"),
                estudiante.get("correo"),
                estudiante.get("programa"),
                estudiante.get("cohorte"),
                estudiante.get("total_evaluaciones"),
                estudiante.get("promedio_nota"),
                estudiante.get("promedio_asistencia"),
                estudiante.get("estado"),
                "Sí" if estudiante.get("es_certificable") else "No",
                estudiante.get("motivo_inconsistencia") or "",
            ]
        )

    _escribir_tabla(resultados, encabezados_resultados, filas_resultados)

    # ------------------------- Hoja Inconsistencias -------------------------
    inconsistencias = libro.create_sheet("Inconsistencias")
    encabezados_inconsistencias = [
        "Tipo",
        "Identificación",
        "Nombre",
        "Programa",
        "Descripción",
        "Promedio nota",
        "Asistencia (%)",
    ]

    filas_inconsistencias = []
    for incidencia in datos["inconsistencias"]:
        filas_inconsistencias.append(
            [
                incidencia.get("tipo"),
                incidencia.get("identificacion"),
                incidencia.get("nombre"),
                incidencia.get("programa"),
                incidencia.get("descripcion"),
                incidencia.get("promedio_nota"),
                incidencia.get("promedio_asistencia"),
            ]
        )

    _escribir_tabla(
        inconsistencias, encabezados_inconsistencias, filas_inconsistencias
    )

    # Guardo el libro en la caja de memoria y vuelvo al inicio para enviarlo.
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer
